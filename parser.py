"""Демо-парсер каталога books.toscrape.com → Excel.

Собирает все книги со всех страниц: название, цену, наличие, рейтинг, категорию и ссылку.
Запуск: python parser.py [--pages N] [--out books.xlsx]
"""

import argparse
import logging
import time
from dataclasses import astuple, dataclass
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

log = logging.getLogger("parser")


@dataclass
class Book:
    title: str
    price: float
    in_stock: int
    rating: int
    category: str
    url: str


HEADERS = ("Название", "Цена, £", "В наличии, шт", "Рейтинг", "Категория", "Ссылка")


class Parser:
    def __init__(self, delay: float = 0.3, retries: int = 3):
        self.session = requests.Session()
        self.session.headers["User-Agent"] = "Mozilla/5.0 (parser-demo)"
        self.delay = delay
        self.retries = retries

    def get(self, url: str) -> BeautifulSoup:
        for attempt in range(1, self.retries + 1):
            try:
                response = self.session.get(url, timeout=15)
                response.raise_for_status()
                response.encoding = "utf-8"
                time.sleep(self.delay)
                return BeautifulSoup(response.text, "html.parser")
            except requests.RequestException as e:
                log.warning("Ошибка %s (попытка %d/%d): %s", url, attempt, self.retries, e)
                time.sleep(attempt * 2)
        raise RuntimeError(f"Не удалось загрузить {url}")

    def book_links(self, max_pages: int | None):
        url, page = urljoin(BASE_URL, "catalogue/page-1.html"), 1
        while url and (max_pages is None or page <= max_pages):
            soup = self.get(url)
            for a in soup.select("article.product_pod h3 a"):
                yield urljoin(url, a["href"])
            next_link = soup.select_one("li.next a")
            url = urljoin(url, next_link["href"]) if next_link else None
            page += 1

    def parse_book(self, url: str) -> Book:
        soup = self.get(url)
        main = soup.select_one("div.product_main")
        stock_text = main.select_one("p.availability").get_text(strip=True)
        rating_class = main.select_one("p.star-rating")["class"][1]
        return Book(
            title=main.h1.get_text(strip=True),
            price=float(main.select_one("p.price_color").get_text(strip=True).lstrip("£Â")),
            in_stock=int("".join(ch for ch in stock_text if ch.isdigit()) or 0),
            rating=RATINGS.get(rating_class, 0),
            category=soup.select("ul.breadcrumb li a")[-1].get_text(strip=True),
            url=url,
        )

    def run(self, max_pages: int | None) -> list[Book]:
        books, seen = [], set()
        for url in self.book_links(max_pages):
            if url in seen:
                continue
            seen.add(url)
            try:
                books.append(self.parse_book(url))
            except (RuntimeError, AttributeError, IndexError, ValueError) as e:
                log.error("Пропускаю %s: %s", url, e)
            if len(books) % 20 == 0:
                log.info("Собрано %d", len(books))
        return books


def save_xlsx(books: list[Book], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Книги"
    ws.append(HEADERS)
    for book in books:
        ws.append(astuple(book))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column, width in zip(range(1, len(HEADERS) + 1), (50, 10, 14, 10, 18, 60)):
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].hyperlink = row[0].value
        row[0].style = "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s: %(message)s")
    args = argparse.ArgumentParser(description=__doc__)
    args.add_argument("--pages", type=int, default=None, help="сколько страниц каталога (по умолчанию все 50)")
    args.add_argument("--out", default="books.xlsx", help="файл Excel для результата")
    opts = args.parse_args()

    books = Parser().run(opts.pages)
    save_xlsx(books, opts.out)
    log.info("Готово: %d книг → %s", len(books), opts.out)


if __name__ == "__main__":
    main()
